import streamlit as st
import fitz  # PyMuPDF
import docx
import os
import json
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from fpdf import FPDF
from google import genai
from google.genai import types
from dotenv import load_dotenv
import ai_optimizer

load_dotenv()

import string

# ─── Page Config (must be first) ───────────────────────────────────────────────
st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─── Premium CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&display=swap');

/* ── Root Variables ── */
:root {
    --bg-primary: #0d1117;
    --bg-secondary: #161b22;
    --bg-card: #1c2333;
    --bg-card-hover: #222d42;
    --accent-blue: #58a6ff;
    --accent-purple: #bc8cff;
    --accent-green: #3fb950;
    --accent-orange: #f78166;
    --accent-yellow: #e3b341;
    --text-primary: #e6edf3;
    --text-secondary: #8b949e;
    --border-color: #30363d;
    --gradient-blue: linear-gradient(135deg, #1a6fb3 0%, #0d4a8f 100%);
    --gradient-purple: linear-gradient(135deg, #6e40c9 0%, #4c2889 100%);
    --shadow-lg: 0 8px 32px rgba(0,0,0,0.4);
    --radius: 12px;
    --radius-sm: 8px;
}

/* ── Global Reset ── */
*, *::before, *::after { box-sizing: border-box; }

html, body, [data-testid="stAppViewContainer"], [data-testid="stApp"] {
    background-color: var(--bg-primary) !important;
    color: var(--text-primary) !important;
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif !important;
}

/* ── Hide Default Streamlit Decoration ── */
#MainMenu, footer, [data-testid="stDeployButton"] { visibility: hidden; }
[data-testid="stDecoration"] { display: none; }
[data-testid="collapsedControl"] { color: var(--text-primary) !important; }

/* ── Scrollbar ── */
::-webkit-scrollbar { width: 6px; }
::-webkit-scrollbar-track { background: var(--bg-secondary); }
::-webkit-scrollbar-thumb { background: var(--border-color); border-radius: 3px; }
::-webkit-scrollbar-thumb:hover { background: var(--accent-blue); }

/* ── Main Content Area ── */
[data-testid="stMain"] {
    background-color: var(--bg-primary) !important;
    padding: 0 !important;
}
[data-testid="stMainBlockContainer"] {
    padding: 2rem 3rem !important;
    max-width: 1400px !important;
}

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background: var(--bg-secondary) !important;
    border-right: 1px solid var(--border-color) !important;
}
[data-testid="stSidebar"] > div:first-child {
    padding: 2rem 1.5rem !important;
}
[data-testid="stSidebarUserContent"] * {
    color: var(--text-primary) !important;
}

/* ── Typography ── */
h1, h2, h3, h4 {
    font-family: 'Inter', sans-serif !important;
    color: var(--text-primary) !important;
    font-weight: 700 !important;
    letter-spacing: -0.02em !important;
}

/* ── Buttons ── */
[data-testid="stButton"] > button {
    background: linear-gradient(135deg, #1f6feb 0%, #1158c7 100%) !important;
    color: white !important;
    border: none !important;
    border-radius: var(--radius-sm) !important;
    padding: 0.6rem 1.4rem !important;
    font-weight: 600 !important;
    font-size: 0.9rem !important;
    letter-spacing: 0.01em !important;
    transition: all 0.2s ease !important;
    box-shadow: 0 2px 8px rgba(31,111,235,0.3) !important;
}
[data-testid="stButton"] > button:hover {
    transform: translateY(-1px) !important;
    box-shadow: 0 4px 16px rgba(31,111,235,0.5) !important;
    background: linear-gradient(135deg, #388bfd 0%, #1f6feb 100%) !important;
}
[data-testid="stButton"] > button:active {
    transform: translateY(0) !important;
}

/* ── Download Button ── */
[data-testid="stDownloadButton"] > button {
    background: linear-gradient(135deg, #238636 0%, #196c2e 100%) !important;
    color: white !important;
    border: none !important;
    border-radius: var(--radius-sm) !important;
    padding: 0.6rem 1.4rem !important;
    font-weight: 600 !important;
    box-shadow: 0 2px 8px rgba(35,134,54,0.3) !important;
    transition: all 0.2s ease !important;
}
[data-testid="stDownloadButton"] > button:hover {
    transform: translateY(-1px) !important;
    box-shadow: 0 4px 16px rgba(35,134,54,0.5) !important;
}

/* ── File Uploader ── */
[data-testid="stFileUploader"] {
    background: var(--bg-card) !important;
    border: 2px dashed var(--border-color) !important;
    border-radius: var(--radius) !important;
    transition: border-color 0.2s ease !important;
}
[data-testid="stFileUploader"]:hover {
    border-color: var(--accent-blue) !important;
}
[data-testid="stFileUploaderDropzoneInput"] + div {
    color: var(--text-secondary) !important;
}

/* ── Metric Cards ── */
[data-testid="stMetric"] {
    background: var(--bg-card) !important;
    border: 1px solid var(--border-color) !important;
    border-radius: var(--radius) !important;
    padding: 1.2rem 1.5rem !important;
    transition: all 0.2s ease !important;
}
[data-testid="stMetric"]:hover {
    border-color: var(--accent-blue) !important;
    box-shadow: 0 4px 16px rgba(88,166,255,0.1) !important;
    transform: translateY(-2px) !important;
}
[data-testid="stMetricLabel"] {
    color: var(--text-secondary) !important;
    font-size: 0.8rem !important;
    font-weight: 500 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.06em !important;
}
[data-testid="stMetricValue"] {
    color: var(--text-primary) !important;
    font-size: 2rem !important;
    font-weight: 800 !important;
}
[data-testid="stMetricDelta"] {
    font-weight: 600 !important;
}

/* ── Expander ── */
[data-testid="stExpander"] {
    background: var(--bg-card) !important;
    border: 1px solid var(--border-color) !important;
    border-radius: var(--radius) !important;
    margin-bottom: 0.75rem !important;
    overflow: hidden !important;
    transition: all 0.2s ease !important;
}
[data-testid="stExpander"]:hover {
    border-color: var(--accent-blue) !important;
}
[data-testid="stExpanderToggleIcon"] { color: var(--accent-blue) !important; }
summary {
    color: var(--text-primary) !important;
    font-weight: 600 !important;
    padding: 0.9rem 1.2rem !important;
}

/* ── Alert Boxes ── */
[data-testid="stAlert"] {
    border-radius: var(--radius) !important;
    border: none !important;
}
.stSuccess {
    background: rgba(63,185,80,0.1) !important;
    border-left: 4px solid var(--accent-green) !important;
    color: var(--text-primary) !important;
}
.stWarning {
    background: rgba(227,179,65,0.1) !important;
    border-left: 4px solid var(--accent-yellow) !important;
    color: var(--text-primary) !important;
}
.stError {
    background: rgba(247,129,102,0.1) !important;
    border-left: 4px solid var(--accent-orange) !important;
    color: var(--text-primary) !important;
}
.stInfo {
    background: rgba(88,166,255,0.08) !important;
    border-left: 4px solid var(--accent-blue) !important;
    color: var(--text-primary) !important;
}

/* ── Spinner ── */
[data-testid="stSpinner"] { color: var(--accent-blue) !important; }

/* ── Slider ── */
[data-testid="stSlider"] [role="slider"] {
    background: var(--accent-blue) !important;
}
[data-testid="stSlider"] [data-baseweb="slider"] > div > div:first-child {
    background: linear-gradient(135deg, #1f6feb, #58a6ff) !important;
}

/* ── Chat ── */
[data-testid="stChatMessage"] {
    background: var(--bg-card) !important;
    border: 1px solid var(--border-color) !important;
    border-radius: var(--radius) !important;
    margin-bottom: 0.75rem !important;
}
[data-testid="stChatInputContainer"] {
    background: var(--bg-secondary) !important;
    border: 1px solid var(--border-color) !important;
    border-radius: var(--radius) !important;
}
[data-testid="stChatInputTextArea"] {
    background: transparent !important;
    color: var(--text-primary) !important;
}

/* ── Input Overrides ── */
[data-baseweb="input"], [data-baseweb="textarea"] {
    background-color: var(--bg-card) !important;
    border-color: var(--border-color) !important;
    color: var(--text-primary) !important;
    border-radius: var(--radius-sm) !important;
}
[data-baseweb="select"] > div {
    background-color: var(--bg-card) !important;
    border-color: var(--border-color) !important;
    color: var(--text-primary) !important;
}

/* ── Divider ── */
hr {
    border: none !important;
    border-top: 1px solid var(--border-color) !important;
    margin: 2.5rem 0 !important;
}

/* ── Columns ── */
[data-testid="stHorizontalBlock"] { gap: 1.5rem !important; }

/* ── Plotly charts dark background ── */
.js-plotly-plot .plotly {
    border-radius: var(--radius) !important;
}

/* ── Custom Components ── */
.hero-header {
    background: linear-gradient(135deg, #0d1b3e 0%, #1a2744 50%, #0f1e3d 100%);
    border: 1px solid rgba(88,166,255,0.2);
    border-radius: 16px;
    padding: 2.5rem 3rem;
    margin-bottom: 2rem;
    position: relative;
    overflow: hidden;
}
.hero-header::before {
    content: '';
    position: absolute;
    top: -50%;
    left: -50%;
    width: 200%;
    height: 200%;
    background: radial-gradient(ellipse at center, rgba(88,166,255,0.06) 0%, transparent 60%);
    pointer-events: none;
}
.hero-title {
    font-size: 2.4rem;
    font-weight: 800;
    background: linear-gradient(135deg, #58a6ff 0%, #bc8cff 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    margin: 0 0 0.5rem 0;
    line-height: 1.2;
}
.hero-subtitle {
    color: #8b949e;
    font-size: 1rem;
    font-weight: 400;
    margin: 0;
}

.section-header {
    display: flex;
    align-items: center;
    gap: 0.75rem;
    margin-bottom: 1.25rem;
    padding-bottom: 0.75rem;
    border-bottom: 1px solid var(--border-color);
}
.section-icon {
    font-size: 1.4rem;
    line-height: 1;
}
.section-title {
    font-size: 1.2rem;
    font-weight: 700;
    color: var(--text-primary);
    margin: 0;
}

.ats-score-card {
    background: linear-gradient(135deg, #0d2137 0%, #0a192f 100%);
    border: 1px solid rgba(88,166,255,0.3);
    border-radius: 16px;
    padding: 2rem;
    text-align: center;
    margin-bottom: 1.5rem;
    position: relative;
    overflow: hidden;
}
.ats-score-card::after {
    content: '';
    position: absolute;
    inset: 0;
    background: radial-gradient(ellipse at 50% 0%, rgba(88,166,255,0.08) 0%, transparent 70%);
    pointer-events: none;
}
.ats-score-label {
    color: var(--text-secondary);
    font-size: 0.8rem;
    font-weight: 600;
    text-transform: uppercase;
    letter-spacing: 0.1em;
    margin-bottom: 0.5rem;
}
.ats-score-value {
    font-size: 4rem;
    font-weight: 900;
    background: linear-gradient(135deg, #58a6ff 0%, #79c0ff 100%);
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    background-clip: text;
    line-height: 1;
    margin: 0;
}
.ats-score-pct {
    font-size: 2rem;
    font-weight: 700;
}

.keyword-pill {
    display: inline-block;
    padding: 0.3rem 0.8rem;
    border-radius: 20px;
    font-size: 0.8rem;
    font-weight: 500;
    margin: 0.2rem;
}
.pill-found {
    background: rgba(63,185,80,0.15);
    color: #3fb950;
    border: 1px solid rgba(63,185,80,0.3);
}
.pill-missing {
    background: rgba(247,129,102,0.12);
    color: #f78166;
    border: 1px solid rgba(247,129,102,0.25);
}

.info-card {
    background: var(--bg-card);
    border: 1px solid var(--border-color);
    border-radius: var(--radius);
    padding: 1.5rem;
    margin-bottom: 1rem;
    transition: all 0.2s ease;
}
.info-card:hover {
    border-color: rgba(88,166,255,0.4);
    box-shadow: 0 4px 20px rgba(88,166,255,0.08);
    transform: translateY(-2px);
}

.bullet-item {
    display: flex;
    align-items: flex-start;
    gap: 0.6rem;
    padding: 0.6rem 0;
    border-bottom: 1px solid rgba(48,54,61,0.5);
    color: var(--text-primary);
    font-size: 0.9rem;
    line-height: 1.5;
}
.bullet-item:last-child { border-bottom: none; }
.bullet-dot {
    width: 6px;
    height: 6px;
    border-radius: 50%;
    background: var(--accent-blue);
    margin-top: 0.45rem;
    flex-shrink: 0;
}
.bullet-dot-green { background: var(--accent-green); }
.bullet-dot-red { background: var(--accent-orange); }

.sidebar-logo {
    display: flex;
    align-items: center;
    gap: 0.75rem;
    margin-bottom: 1.75rem;
    padding-bottom: 1.25rem;
    border-bottom: 1px solid var(--border-color);
}
.sidebar-logo-icon { font-size: 2rem; }
.sidebar-logo-text {
    font-size: 1rem;
    font-weight: 700;
    color: var(--text-primary);
}
.sidebar-logo-sub {
    font-size: 0.75rem;
    color: var(--text-secondary);
}

.status-badge {
    display: inline-flex;
    align-items: center;
    gap: 0.4rem;
    padding: 0.3rem 0.75rem;
    border-radius: 20px;
    font-size: 0.8rem;
    font-weight: 600;
}
.badge-success {
    background: rgba(63,185,80,0.15);
    color: #3fb950;
    border: 1px solid rgba(63,185,80,0.3);
}
.badge-info {
    background: rgba(88,166,255,0.12);
    color: #58a6ff;
    border: 1px solid rgba(88,166,255,0.25);
}

.step-history-item {
    display: flex;
    align-items: center;
    gap: 0.8rem;
    padding: 0.65rem 0.9rem;
    border-radius: var(--radius-sm);
    background: rgba(88,166,255,0.05);
    border: 1px solid rgba(88,166,255,0.1);
    margin-bottom: 0.5rem;
    font-size: 0.85rem;
    color: var(--text-primary);
}

/* Stacked label colors */
[data-testid="stSidebar"] label,
[data-testid="stSidebar"] p,
[data-testid="stSidebar"] span {
    color: var(--text-secondary) !important;
}
[data-testid="stSidebar"] h1,
[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 {
    color: var(--text-primary) !important;
}

/* ── Selectbox, radio ── */
[data-baseweb="radio"] span { color: var(--text-primary) !important; }

/* ── Override p tags ── */
p { color: var(--text-primary) !important; }
small { color: var(--text-secondary) !important; }

/* ── Tab bar ── */
[data-testid="stTabs"] [role="tablist"] {
    background: var(--bg-secondary) !important;
    border-radius: var(--radius) !important;
    padding: 0.3rem !important;
    border: 1px solid var(--border-color) !important;
    gap: 0.25rem !important;
}
[data-testid="stTabs"] [role="tab"] {
    border-radius: var(--radius-sm) !important;
    color: var(--text-secondary) !important;
    font-weight: 500 !important;
    padding: 0.5rem 1rem !important;
    transition: all 0.2s ease !important;
}
[data-testid="stTabs"] [role="tab"][aria-selected="true"] {
    background: var(--bg-card) !important;
    color: var(--text-primary) !important;
    font-weight: 600 !important;
}
[data-testid="stTabs"] [role="tabpanel"] {
    padding-top: 1.5rem !important;
}
</style>
""", unsafe_allow_html=True)


# ─── Helpers ───────────────────────────────────────────────────────────────────

def sanitize_for_pdf(text):
    if not text:
        return ""
    text = str(text)
    replacements = {
        '\u201c': '"', '\u201d': '"', '\u2018': "'", '\u2019': "'",
        '\u2013': '-', '\u2014': '-', '\u2026': '...',
        '\u2022': '-', '\u25cf': '-', '\u2605': '*', '\u2606': '*',
        '\t': ' ', '\r': '', '\n': ' '
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    valid_chars = set(string.printable)
    text = ''.join(c for c in text if c in valid_chars)
    words = text.split()
    return " ".join(w[:90] for w in words)


def create_pdf_report(analysis_dict):
    pdf = FPDF()
    pdf.add_page()
    L = pdf.l_margin        # left margin (~10 mm)
    W = pdf.w - 2 * L      # usable page width

    def write_heading(text, bold=False, size=12):
        """Write a heading line, always starting at left margin."""
        pdf.set_x(L)
        pdf.set_font("Arial", 'B' if bold else '', size)
        pdf.cell(W, 10, sanitize_for_pdf(text))
        pdf.ln(9)

    def write_body(text, size=11):
        """Write wrapping body text, always starting at left margin."""
        pdf.set_x(L)
        pdf.set_font("Arial", '', size)
        pdf.multi_cell(W, 8, sanitize_for_pdf(str(text)))

    write_heading("Resume Analysis Report", bold=True, size=16)
    pdf.ln(2)
    write_heading(f"ATS Match: {analysis_dict.get('ats_match_percentage', 0)}%", bold=True, size=13)
    pdf.ln(2)

    assessment = analysis_dict.get('assessment', '')
    write_body(f"Assessment: {assessment}")
    pdf.ln(4)

    write_heading("Strengths:", bold=True)
    for s in analysis_dict.get('strengths', []):
        write_body(f"  - {s}")

    pdf.ln(4)
    write_heading("Weaknesses:", bold=True)
    for w in analysis_dict.get('weaknesses', []):
        write_body(f"  - {w}")

    suggestions = analysis_dict.get('actionable_suggestions', [])
    if suggestions:
        pdf.ln(4)
        write_heading("Actionable Suggestions:", bold=True)
        for sg in suggestions:
            write_body(f"  - {sg}")

    out = pdf.output()
    return bytes(out) if not isinstance(out, bytes) else out


# ─── API Configuration ──────────────────────────────────────────────────────────
import time

api_key = os.environ.get("GEMINI_API_KEY")
client = genai.Client(api_key=api_key) if api_key else None

# Model preference order — falls back automatically on 503 overload
_MODEL_PRIORITY = ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-2.0-flash-lite"]

def _call_gemini(prompt: str, json_mode: bool = True, retries: int = 3) -> str:
    """Call Gemini with automatic retries (exponential backoff) and model fallback."""
    cfg = types.GenerateContentConfig(response_mime_type="application/json") if json_mode else None
    last_err = None
    for model in _MODEL_PRIORITY:
        for attempt in range(retries):
            try:
                resp = client.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=cfg,
                )
                return resp.text
            except Exception as e:
                last_err = e
                msg = str(e)
                if "503" in msg or "UNAVAILABLE" in msg:
                    wait = 2 ** attempt          # 1s, 2s, 4s
                    time.sleep(wait)
                    continue                     # retry same model
                elif "429" in msg or "RESOURCE_EXHAUSTED" in msg:
                    time.sleep(5)
                    continue
                else:
                    break                        # non-retryable — try next model
    raise last_err


# ─── File Readers ───────────────────────────────────────────────────────────────

def read_pdf(file):
    text = ""
    pdf = fitz.open(stream=file.read(), filetype="pdf")
    for page in pdf:
        text += page.get_text()
    return text


def read_docx(file):
    doc = docx.Document(file)
    return "\n".join([p.text for p in doc.paragraphs])


def extract_text(uploaded_file):
    if uploaded_file.name.endswith(".pdf"):
        return read_pdf(uploaded_file)
    elif uploaded_file.name.endswith(".docx"):
        return read_docx(uploaded_file)
    return "Unsupported file format. Please upload a .pdf or .docx file."


# ─── AI Calls ───────────────────────────────────────────────────────────────────

def get_job_recommendations(resume_text):
    prompt = f"""
    You are a career advisor. Based on the following resume, suggest 3 to 5 job roles that best match this candidate's skills and experience.
    Resume:
    {resume_text}

    Return ONLY a valid JSON array (no markdown, no backticks) of objects with these exact keys:
    [
      {{
        "title": "<job title>",
        "why_it_fits": "<1-2 sentence explanation>",
        "where_to_apply": "<suggested platforms or company types e.g. LinkedIn, startups, FAANG>"
      }}
    ]
    """
    return json.loads(_call_gemini(prompt))


def get_improvement_tips(resume_text, job_description):
    prompt = f"""
    You are an expert resume coach. Analyze the resume below and give specific, actionable improvement advice.
    Resume:
    {resume_text}
    Job Description:
    {job_description}

    Return ONLY a valid JSON object (no markdown, no backticks) with these exact keys:
    {{
      "bullet_rewrites": ["<original bullet> → <improved version>"],
      "skills_to_add": ["skill or certification to pursue"],
      "formatting_tips": ["formatting or structure tip"]
    }}
    """
    return json.loads(_call_gemini(prompt))


def analyze_resume(resume_text, job_description):
    prompt = f"""
    You are a resume evaluation expert. Analyze the following resume:
    {resume_text}
    in the context of this job description:
    {job_description}
    
    Return ONLY a valid JSON object (no markdown formatting, no backticks) with the following exact keys:
    {{
      "ats_match_percentage": <int 0-100 indicating overall alignment>,
      "technical_score": <int 0-100>,
      "leadership_score": <int 0-100>,
      "communication_score": <int 0-100>,
      "found_keywords": ["provide a comprehensive list of at least 15-20 labels found including technical skills, tools, and industry terms (Exclude generic buzzwords like 'motivated', 'detail-oriented', 'hard-working')"],
      "missing_keywords": ["provide an exhaustive list of at least 15-20 labels missing from the resume but required by the JD (Exclude generic buzzwords like 'motivated', 'detail-oriented', 'hard-working')"],
      "assessment": "<short overall assessment string>",
      "strengths": ["strength1", "strength2"],
      "weaknesses": ["weakness1", "weakness2"],
      "actionable_suggestions": ["suggestion1"]
    }}
    """
    return json.loads(_call_gemini(prompt))


# ─── Sidebar ────────────────────────────────────────────────────────────────────

with st.sidebar:
    st.markdown("""
        <div class="sidebar-logo">
            <div class="sidebar-logo-icon">📄</div>
            <div>
                <div class="sidebar-logo-text">ResumeAI</div>
                <div class="sidebar-logo-sub">Powered by Gemini 2.5</div>
            </div>
        </div>
    """, unsafe_allow_html=True)

    st.markdown("### Upload Documents")
    st.markdown("<small>Accepts PDF and DOCX formats</small>", unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)

    resume_file = st.file_uploader(
        "📋 Resume",
        type=["pdf", "docx"],
        help="Upload your resume in PDF or DOCX format"
    )
    st.markdown("<br>", unsafe_allow_html=True)
    jd_file = st.file_uploader(
        "💼 Job Description",
        type=["pdf", "docx"],
        help="Upload the target job description"
    )

    st.markdown("<br>", unsafe_allow_html=True)

    # File status indicators
    if resume_file:
        st.markdown('<span class="status-badge badge-success">✓ Resume uploaded</span>', unsafe_allow_html=True)
    if jd_file:
        st.markdown('<span class="status-badge badge-success">✓ JD uploaded</span>', unsafe_allow_html=True)

    st.markdown("<br><br>", unsafe_allow_html=True)
    st.markdown("---")
    st.markdown("""
        <div style="color: #8b949e; font-size: 0.78rem; line-height: 1.6;">
            <strong style="color: #e6edf3;">How it works</strong><br><br>
            1. Upload your resume<br>
            2. Upload a job description<br>
            3. Click Analyze<br>
            4. Review your ATS score,<br>&nbsp;&nbsp;&nbsp;keywords, and insights<br>
            5. Optimize & export PDF
        </div>
    """, unsafe_allow_html=True)


# ─── Hero Header ────────────────────────────────────────────────────────────────

st.markdown("""
<div class="hero-header">
    <div class="hero-title">🔍 AI Resume Analyzer</div>
    <div class="hero-subtitle">
        Instantly evaluate your resume against any job description — get ATS scores, keyword gaps, and actionable improvements powered by Google Gemini.
    </div>
</div>
""", unsafe_allow_html=True)

# ─── Analyze Button ──────────────────────────────────────────────────────────────

col_btn, col_hint = st.columns([1, 3])
with col_btn:
    analyze_button = st.button("🚀 Analyze Resume", use_container_width=True)
with col_hint:
    if not resume_file or not jd_file:
        st.markdown('<p style="color:#8b949e; margin-top:0.65rem; font-size:0.9rem;">← Upload both files in the sidebar to get started</p>', unsafe_allow_html=True)

if analyze_button and resume_file and jd_file:
    api_key_val = api_key and api_key.strip()
    if not api_key_val:
        st.error("⚠️ The GEMINI_API_KEY environment variable is not set. Please restart the app with this variable set.")
    elif not api_key_val.startswith("AIza"):
        st.error("⚠️ The GEMINI_API_KEY doesn't look like a valid Google Gemini key (should start with 'AIza').")
    else:
        with st.spinner("🔍 Analyzing your resume with Gemini AI…"):
            try:
                resume_text = extract_text(resume_file)
                job_text = extract_text(jd_file)
                st.session_state.resume_text_cache = resume_text
                st.session_state.job_text_cache = job_text
                analysis = analyze_resume(resume_text, job_text)
                st.session_state.last_analysis = analysis
                # Clear stale sub-results when re-analyzing
                for key in ["job_recommendations", "improvement_tips", "optimization_results", "chat_history"]:
                    st.session_state.pop(key, None)
            except Exception as e:
                st.error(f"⚠️ {str(e)}")

elif analyze_button:
    st.warning("⚠️ Please upload both a resume and a job description before analyzing.")


# ─── Results ────────────────────────────────────────────────────────────────────

if "last_analysis" in st.session_state:
    analysis = st.session_state.last_analysis
    ats_score = analysis.get('ats_match_percentage', 0)

    st.markdown("<br>", unsafe_allow_html=True)

    # ── ATS Score Banner ──
    score_color = "#3fb950" if ats_score >= 75 else ("#e3b341" if ats_score >= 50 else "#f78166")
    st.markdown(f"""
    <div class="ats-score-card">
        <div class="ats-score-label">Overall ATS Match Score</div>
        <div class="ats-score-value" style="background: linear-gradient(135deg, {score_color} 0%, #58a6ff 100%); -webkit-background-clip: text; -webkit-text-fill-color: transparent; background-clip: text;">
            {ats_score}<span class="ats-score-pct">%</span>
        </div>
        <div style="color: #8b949e; font-size: 0.85rem; margin-top: 0.5rem;">
            {"🟢 Strong match — you're competitive for this role" if ats_score >= 75 else ("🟡 Moderate match — some gaps to address" if ats_score >= 50 else "🔴 Low match — significant improvements needed")}
        </div>
    </div>
    """, unsafe_allow_html=True)

    # ── Skill Score Metrics ──
    m1, m2, m3 = st.columns(3)
    with m1:
        st.metric("🛠️ Technical Score", f"{analysis.get('technical_score', 0)}%")
    with m2:
        st.metric("🏆 Leadership Score", f"{analysis.get('leadership_score', 0)}%")
    with m3:
        st.metric("💬 Communication Score", f"{analysis.get('communication_score', 0)}%")

    st.markdown("<br>", unsafe_allow_html=True)

    # ── Tabs ──
    tab1, tab2, tab3, tab4, tab5 = st.tabs([
        "📊 Analytics", "🔑 Keywords", "📝 Assessment",
        "💼 Opportunities", "🚀 Optimizer"
    ])

    # ────── Tab 1: Analytics ──────
    with tab1:
        col1, col2 = st.columns([1, 1], gap="large")

        with col1:
            st.markdown('<div class="section-header"><span class="section-icon">📡</span><span class="section-title">Skills Radar</span></div>', unsafe_allow_html=True)
            radar_df = pd.DataFrame(dict(
                r=[analysis.get('technical_score', 0),
                   analysis.get('leadership_score', 0),
                   analysis.get('communication_score', 0)],
                theta=['Technical', 'Leadership', 'Communication']
            ))
            fig_radar = px.line_polar(radar_df, r='r', theta='theta', line_close=True)
            fig_radar.update_traces(
                fill='toself',
                fillcolor='rgba(88,166,255,0.15)',
                line=dict(color='#58a6ff', width=2.5),
            )
            fig_radar.update_layout(
                paper_bgcolor='rgba(0,0,0,0)',
                plot_bgcolor='rgba(0,0,0,0)',
                polar=dict(
                    bgcolor='rgba(22,27,34,0.8)',
                    radialaxis=dict(visible=True, range=[0, 100], color='#8b949e', gridcolor='#30363d'),
                    angularaxis=dict(color='#e6edf3', gridcolor='#30363d')
                ),
                font=dict(family='Inter', color='#e6edf3'),
                margin=dict(l=30, r=30, t=30, b=30),
                height=320,
            )
            st.plotly_chart(fig_radar, use_container_width=True, config={'displayModeBar': False})

        with col2:
            st.markdown('<div class="section-header"><span class="section-icon">📈</span><span class="section-title">Score Breakdown</span></div>', unsafe_allow_html=True)
            scores = {
                'Technical': analysis.get('technical_score', 0),
                'Leadership': analysis.get('leadership_score', 0),
                'Communication': analysis.get('communication_score', 0),
                'Overall ATS': ats_score,
            }
            colors = ['#58a6ff', '#bc8cff', '#3fb950', '#e3b341']
            fig_bar = go.Figure(go.Bar(
                x=list(scores.values()),
                y=list(scores.keys()),
                orientation='h',
                marker=dict(color=colors, line=dict(width=0)),
                text=[f"{v}%" for v in scores.values()],
                textposition='outside',
                textfont=dict(color='#e6edf3', size=13, family='Inter'),
            ))
            fig_bar.update_layout(
                paper_bgcolor='rgba(0,0,0,0)',
                plot_bgcolor='rgba(0,0,0,0)',
                xaxis=dict(range=[0, 115], gridcolor='#30363d', tickfont=dict(color='#8b949e')),
                yaxis=dict(gridcolor='#30363d', tickfont=dict(color='#e6edf3', size=13)),
                font=dict(family='Inter', color='#e6edf3'),
                margin=dict(l=10, r=40, t=20, b=20),
                height=280,
                bargap=0.35,
            )
            st.plotly_chart(fig_bar, use_container_width=True, config={'displayModeBar': False})

    # ────── Tab 2: Keywords ──────
    with tab2:
        col_kw1, col_kw2 = st.columns(2, gap="large")

        with col_kw1:
            st.markdown('<div class="section-header"><span class="section-icon">✅</span><span class="section-title">Found Keywords</span></div>', unsafe_allow_html=True)
            found = analysis.get('found_keywords', [])
            if found:
                pills_html = "".join(f'<span class="keyword-pill pill-found">{k}</span>' for k in found)
                st.markdown(f'<div style="line-height:2.2;">{pills_html}</div>', unsafe_allow_html=True)
                st.markdown(f'<p style="color:#8b949e; font-size:0.8rem; margin-top:0.75rem;">{len(found)} keywords matched</p>', unsafe_allow_html=True)
            else:
                st.markdown('<p style="color:#8b949e;">No keywords found.</p>', unsafe_allow_html=True)

        with col_kw2:
            st.markdown('<div class="section-header"><span class="section-icon">❌</span><span class="section-title">Missing Keywords</span></div>', unsafe_allow_html=True)
            missing = analysis.get('missing_keywords', [])
            if missing:
                pills_html = "".join(f'<span class="keyword-pill pill-missing">{k}</span>' for k in missing)
                st.markdown(f'<div style="line-height:2.2;">{pills_html}</div>', unsafe_allow_html=True)
                st.markdown(f'<p style="color:#8b949e; font-size:0.8rem; margin-top:0.75rem;">{len(missing)} keywords to add</p>', unsafe_allow_html=True)
            else:
                st.markdown('<p style="color:#8b949e;">No missing keywords.</p>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)

        # Keyword Coverage Donut
        if found or missing:
            st.markdown('<div class="section-header"><span class="section-icon">🍩</span><span class="section-title">Keyword Coverage</span></div>', unsafe_allow_html=True)
            fig_donut = go.Figure(go.Pie(
                labels=['Found', 'Missing'],
                values=[len(found), len(missing)],
                hole=0.65,
                marker=dict(colors=['#3fb950', '#f78166'], line=dict(width=0)),
                textinfo='label+percent',
                textfont=dict(color='#e6edf3', family='Inter', size=13),
            ))
            fig_donut.update_layout(
                paper_bgcolor='rgba(0,0,0,0)',
                plot_bgcolor='rgba(0,0,0,0)',
                font=dict(family='Inter', color='#e6edf3'),
                showlegend=False,
                margin=dict(l=20, r=20, t=10, b=10),
                height=260,
                annotations=[dict(
                    text=f"<b>{len(found)}/{len(found)+len(missing)}</b>",
                    x=0.5, y=0.5, font_size=20, font_color='#e6edf3',
                    showarrow=False, font_family='Inter'
                )]
            )
            col_donut, _ = st.columns([1, 1])
            with col_donut:
                st.plotly_chart(fig_donut, use_container_width=True, config={'displayModeBar': False})

    # ────── Tab 3: Assessment ──────
    with tab3:
        # Overall Assessment
        st.markdown('<div class="section-header"><span class="section-icon">📝</span><span class="section-title">Overall Assessment</span></div>', unsafe_allow_html=True)
        st.markdown(f'<div class="info-card"><p style="font-size:1rem; line-height:1.7; margin:0; color:#e6edf3;">{analysis.get("assessment", "")}</p></div>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        col_sw1, col_sw2 = st.columns(2, gap="large")

        with col_sw1:
            st.markdown('<div class="section-header"><span class="section-icon">🌟</span><span class="section-title">Strengths</span></div>', unsafe_allow_html=True)
            strengths_html = ''.join(f'<div class="bullet-item"><span class="bullet-dot bullet-dot-green"></span>{s}</div>' for s in analysis.get('strengths', []))
            st.markdown(f'<div class="info-card">{strengths_html}</div>', unsafe_allow_html=True)

        with col_sw2:
            st.markdown('<div class="section-header"><span class="section-icon">⚠️</span><span class="section-title">Areas to Improve</span></div>', unsafe_allow_html=True)
            weaknesses_html = ''.join(f'<div class="bullet-item"><span class="bullet-dot bullet-dot-red"></span>{w}</div>' for w in analysis.get('weaknesses', []))
            st.markdown(f'<div class="info-card">{weaknesses_html}</div>', unsafe_allow_html=True)

        # Actionable Suggestions
        suggestions = analysis.get('actionable_suggestions', [])
        if suggestions:
            st.markdown("<br>", unsafe_allow_html=True)
            st.markdown('<div class="section-header"><span class="section-icon">💡</span><span class="section-title">Actionable Suggestions</span></div>', unsafe_allow_html=True)
            suggestions_html = ''.join(f'<div class="bullet-item"><span class="bullet-dot"></span>{sg}</div>' for sg in suggestions)
            st.markdown(f'<div class="info-card">{suggestions_html}</div>', unsafe_allow_html=True)

        st.markdown("<br>", unsafe_allow_html=True)
        pdf_bytes = create_pdf_report(analysis)
        st.download_button(
            label="📥 Export Full Report as PDF",
            data=pdf_bytes,
            file_name="resume_analysis.pdf",
            mime="application/pdf",
            use_container_width=False
        )

    # ────── Tab 4: Opportunities ──────
    with tab4:
        col_a, col_b = st.columns([1, 1], gap="large")

        # ── Job Recommendations ──
        with col_a:
            st.markdown('<div class="section-header"><span class="section-icon">💼</span><span class="section-title">Job Role Recommendations</span></div>', unsafe_allow_html=True)
            if "job_recommendations" not in st.session_state:
                st.markdown('<p style="color:#8b949e; font-size:0.9rem; margin-bottom:1rem;">Discover the best-matching job roles based on your resume profile.</p>', unsafe_allow_html=True)
                if st.button("🔍 Find Matching Roles", key="get_recs"):
                    with st.spinner("Analyzing your profile…"):
                        try:
                            recs = get_job_recommendations(st.session_state.resume_text_cache)
                            st.session_state.job_recommendations = recs
                            st.rerun()
                        except Exception as e:
                            st.error(f"Error: {e}")
            else:
                for rec in st.session_state.job_recommendations:
                    with st.expander(f"🏷️ {rec.get('title', 'Role')}"):
                        st.markdown(f'<div class="bullet-item"><span class="bullet-dot bullet-dot-green"></span><strong>Why it fits:</strong>&nbsp;{rec.get("why_it_fits", "")}</div>', unsafe_allow_html=True)
                        st.markdown(f'<div class="bullet-item"><span class="bullet-dot"></span><strong>Where to apply:</strong>&nbsp;{rec.get("where_to_apply", "")}</div>', unsafe_allow_html=True)
                if st.button("🔄 Refresh", key="refresh_recs"):
                    del st.session_state.job_recommendations
                    st.rerun()

        # ── Improvement Coach ──
        with col_b:
            st.markdown('<div class="section-header"><span class="section-icon">🛠️</span><span class="section-title">Resume Improvement Coach</span></div>', unsafe_allow_html=True)
            if "improvement_tips" not in st.session_state:
                st.markdown('<p style="color:#8b949e; font-size:0.9rem; margin-bottom:1rem;">Get tailored bullet rewrites, skills to add, and formatting tips.</p>', unsafe_allow_html=True)
                if st.button("💡 Get Improvement Tips", key="get_tips"):
                    with st.spinner("Generating personalized tips…"):
                        try:
                            tips = get_improvement_tips(
                                st.session_state.resume_text_cache,
                                st.session_state.job_text_cache
                            )
                            st.session_state.improvement_tips = tips
                            st.rerun()
                        except Exception as e:
                            st.error(f"Error: {e}")
            else:
                tips = st.session_state.improvement_tips
                with st.expander("✏️ Bullet Point Rewrites", expanded=True):
                    for b in tips.get("bullet_rewrites", []):
                        st.markdown(f'<div class="bullet-item"><span class="bullet-dot"></span>{b}</div>', unsafe_allow_html=True)
                with st.expander("📚 Skills & Certifications to Add"):
                    for s in tips.get("skills_to_add", []):
                        st.markdown(f'<div class="bullet-item"><span class="bullet-dot bullet-dot-green"></span>{s}</div>', unsafe_allow_html=True)
                with st.expander("📐 Formatting & Structure Tips"):
                    for f_tip in tips.get("formatting_tips", []):
                        st.markdown(f'<div class="bullet-item"><span class="bullet-dot"></span>{f_tip}</div>', unsafe_allow_html=True)
                if st.button("🔄 Refresh Tips", key="refresh_tips"):
                    del st.session_state.improvement_tips
                    st.rerun()

    # ────── Tab 5: Optimizer ──────
    with tab5:
        st.markdown('<div class="section-header"><span class="section-icon">🚀</span><span class="section-title">ATS Score Optimizer (Hill Climbing)</span></div>', unsafe_allow_html=True)
        st.markdown('<div class="info-card"><p style="color:#8b949e; font-size:0.9rem; margin:0;">Iteratively adds high-impact missing keywords and replaces low-weight ones to maximize your ATS match. Configure below and click Optimize.</p></div>', unsafe_allow_html=True)
        st.markdown("<br>", unsafe_allow_html=True)

        col_cfg1, col_cfg2 = st.columns(2)
        with col_cfg1:
            max_k = st.slider("Max Keywords to Add", 1, 15, 5)
        with col_cfg2:
            tgt_s = st.slider("Target ATS Score %", 50, 100, 90)

        st.markdown("<br>", unsafe_allow_html=True)
        if st.button("📈 Optimize Resume Score", key="optimize"):
            with st.spinner("Running Hill Climbing Optimization…"):
                try:
                    found_kw = analysis.get('found_keywords', [])
                    missing_kw = analysis.get('missing_keywords', [])
                    jd_text = st.session_state.job_text_cache
                    results = ai_optimizer.hill_climbing_optimize(found_kw, missing_kw, jd_text, max_extra_keywords=max_k, target_score=tgt_s)
                    st.session_state.optimization_results = results
                    st.rerun()
                except Exception as e:
                    st.error(f"Optimization Error: {e}")

        if "optimization_results" in st.session_state:
            res = st.session_state.optimization_results
            orig = res['original_score']
            opt = res['optimized_score']
            delta = round(opt - orig, 2)

            st.markdown("<br>", unsafe_allow_html=True)
            col_o1, col_o2, col_o3 = st.columns(3)
            with col_o1:
                st.metric("📉 Original Score", f"{orig}%")
            with col_o2:
                st.metric("📈 Optimized Score", f"{opt}%", delta=f"+{delta}%")
            with col_o3:
                st.metric("🎯 Score Gain", f"+{delta}%")

            st.markdown("<br>", unsafe_allow_html=True)
            col_add, col_rem = st.columns(2, gap="large")
            with col_add:
                if res['added']:
                    st.markdown('<div class="section-header"><span class="section-icon">➕</span><span class="section-title">Keywords to Add</span></div>', unsafe_allow_html=True)
                    pills = "".join(f'<span class="keyword-pill pill-found">{k}</span>' for k in res['added'])
                    st.markdown(f'<div style="line-height:2.4;">{pills}</div>', unsafe_allow_html=True)

            with col_rem:
                if res['removed']:
                    st.markdown('<div class="section-header"><span class="section-icon">🔁</span><span class="section-title">Keywords to Replace</span></div>', unsafe_allow_html=True)
                    pills = "".join(f'<span class="keyword-pill pill-missing">{k}</span>' for k in res['removed'])
                    st.markdown(f'<div style="line-height:2.4;">{pills}</div>', unsafe_allow_html=True)

            st.markdown("<br>", unsafe_allow_html=True)
            with st.expander("📋 Step-by-Step Optimization History"):
                for step in res['history']:
                    st.markdown(f'<div class="step-history-item">⚡ {step}</div>', unsafe_allow_html=True)

            if st.button("🔄 Reset Optimization", key="reset_opt"):
                del st.session_state.optimization_results
                st.rerun()

    # ─── Chat ──────────────────────────────────────────────────────────────────
    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("---")
    st.markdown('<div class="section-header"><span class="section-icon">💬</span><span class="section-title">Chat with this Resume</span></div>', unsafe_allow_html=True)
    st.markdown('<p style="color:#8b949e; font-size:0.9rem; margin-top:-0.75rem; margin-bottom:1rem;">Ask anything about the resume — strengths, fit for specific roles, interview prep questions, etc.</p>', unsafe_allow_html=True)

    if "chat_history" not in st.session_state:
        st.session_state.chat_history = []

    for msg in st.session_state.chat_history:
        st.chat_message(msg["role"]).write(msg["text"])

    if user_q := st.chat_input("Ask a question about this resume…"):
        st.chat_message("user").write(user_q)
        st.session_state.chat_history.append({"role": "user", "text": user_q})
        qa_prompt = f"Resume:\n{st.session_state.resume_text_cache}\n\nQuestion: {user_q}"
        try:
            answer_text = _call_gemini(qa_prompt, json_mode=False)
            st.chat_message("assistant").write(answer_text)
            st.session_state.chat_history.append({"role": "assistant", "text": answer_text})
        except Exception as e:
            st.error(f"Chat Error: {e}")

else:
    # ─── Empty State ─────────────────────────────────────────────────────────
    st.markdown("<br><br>", unsafe_allow_html=True)
    st.markdown("""
    <div style="text-align:center; padding: 4rem 2rem;">
        <div style="font-size: 4rem; margin-bottom: 1.5rem;">📄</div>
        <h3 style="color: #e6edf3; font-size: 1.5rem; margin-bottom: 0.75rem;">Ready to analyze your resume?</h3>
        <p style="color: #8b949e; font-size: 1rem; max-width: 480px; margin: 0 auto 2rem auto; line-height: 1.7;">
            Upload your resume and a job description in the sidebar, then click <strong style="color:#58a6ff;">Analyze Resume</strong> to get an instant ATS score, keyword gaps, and AI-powered improvement tips.
        </p>
        <div style="display: flex; justify-content: center; gap: 2rem; flex-wrap: wrap; margin-top: 2rem;">
            <div style="background: #1c2333; border: 1px solid #30363d; border-radius: 12px; padding: 1.5rem; width: 180px;">
                <div style="font-size: 2rem; margin-bottom: 0.5rem;">🎯</div>
                <div style="color: #e6edf3; font-weight: 600; font-size: 0.9rem;">ATS Scoring</div>
                <div style="color: #8b949e; font-size: 0.8rem; margin-top: 0.25rem;">See how well you match</div>
            </div>
            <div style="background: #1c2333; border: 1px solid #30363d; border-radius: 12px; padding: 1.5rem; width: 180px;">
                <div style="font-size: 2rem; margin-bottom: 0.5rem;">🔑</div>
                <div style="color: #e6edf3; font-weight: 600; font-size: 0.9rem;">Keyword Gap Analysis</div>
                <div style="color: #8b949e; font-size: 0.8rem; margin-top: 0.25rem;">Find what's missing</div>
            </div>
            <div style="background: #1c2333; border: 1px solid #30363d; border-radius: 12px; padding: 1.5rem; width: 180px;">
                <div style="font-size: 2rem; margin-bottom: 0.5rem;">🚀</div>
                <div style="color: #e6edf3; font-weight: 600; font-size: 0.9rem;">AI Optimizer</div>
                <div style="color: #8b949e; font-size: 0.8rem; margin-top: 0.25rem;">Maximize your score</div>
            </div>
            <div style="background: #1c2333; border: 1px solid #30363d; border-radius: 12px; padding: 1.5rem; width: 180px;">
                <div style="font-size: 2rem; margin-bottom: 0.5rem;">💬</div>
                <div style="color: #e6edf3; font-weight: 600; font-size: 0.9rem;">Resume Chat</div>
                <div style="color: #8b949e; font-size: 0.8rem; margin-top: 0.25rem;">Ask anything about your CV</div>
            </div>
        </div>
    </div>
    """, unsafe_allow_html=True)
